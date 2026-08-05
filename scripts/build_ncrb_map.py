#!/usr/bin/env python3
"""Build the state-by-year crime map data from NCRB district returns.

Reads:
  data/ncrb/source/District_wise_IPC_crimes_reported_2001_2018.xlsx
  data/ncrb/source/india_districts.geojson

Writes:
  public/pew-infographics/reports/india-crime-map-data.js

Every value is aggregated from the workbook at build time. Nothing is typed in.

Aggregation basis: for each state-year, all district rows are summed, excluding
the sheets' own aggregate rows ("TOTAL" / "ZZ TOTAL") so nothing is counted
twice. That basis is available in every sheet, whereas the aggregate rows are
missing from some, and it runs about 3% above the sheets' own TOTAL rows
because it includes railway/CID/crime-branch reporting units that those rows
leave out. The report states this.

Only 2001-2015 is emitted: the 2016-2018 sheets carry a corrupt YEAR column
(it holds a row index, not a year), so they cannot be attributed to a year
safely and are skipped rather than guessed at.
"""

import json
import math
import os
import re
from collections import defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "ncrb", "source")
XLSX = os.path.join(SRC, "District_wise_IPC_crimes_reported_2001_2018.xlsx")
GEO = os.path.join(SRC, "india_districts.geojson")
OUT = os.path.join(ROOT, "public", "pew-infographics", "reports", "india-crime-map-data.js")

# Sheets whose YEAR column is usable. 2016/2017/2018 are excluded on purpose.
SHEETS = ["2001-2012", "2013", "2014", "2015"]

# Header spellings differ per sheet ("ROBERRY", "KIDNAPPING ABDUCTION", ...).
# Only heads present in EVERY sheet above are emitted, so no series has a
# year-shaped hole in it.
CANON = {
    "MURDER": "murder",
    "RAPE": "rape",
    "RIOTS": "riots",
    "THEFT": "theft",
    "ROBBERY": "robbery",
    "ROBERRY": "robbery",
    "DOWRY DEATHS": "dowry_deaths",
    "DOWRY DEATH": "dowry_deaths",
    "ASSAULT ON WOMEN": "assault_on_women",
    "ASSAULT ON OWMEN": "assault_on_women",
}

LABELS = {
    "murder": "Murder",
    "rape": "Rape",
    "riots": "Riots",
    "theft": "Theft",
    "robbery": "Robbery",
    "dowry_deaths": "Dowry deaths",
    "assault_on_women": "Assault on women",
}

AGGREGATE_ROWS = {"TOTAL", "ZZ TOTAL", "TOTAL STATE(S)", "GRAND TOTAL"}

# NCRB spellings → the GeoJSON's st_nm spellings.
STATE_ALIASES = {
    "A & N ISLANDS": "Andaman and Nicobar Islands",
    "A&N ISLANDS": "Andaman and Nicobar Islands",
    "ANDAMAN & NICOBAR ISLANDS": "Andaman and Nicobar Islands",
    "D & N HAVELI": "Dadra and Nagar Haveli and Daman and Diu",
    "D&N HAVELI": "Dadra and Nagar Haveli and Daman and Diu",
    "DAMAN & DIU": "Dadra and Nagar Haveli and Daman and Diu",
    "DELHI UT": "Delhi",
    "DELHI": "Delhi",
    "ORISSA": "Odisha",
    "ODISHA": "Odisha",
    "PONDICHERRY": "Puducherry",
    "PUDUCHERRY": "Puducherry",
    "JAMMU & KASHMIR": "Jammu and Kashmir",
    "UTTARANCHAL": "Uttarakhand",
    "UTTARAKHAND": "Uttarakhand",
    "CHATISGARH": "Chhattisgarh",
    "CHHATTISGARH": "Chhattisgarh",
    "TELENGANA": "Telangana",
}


def canon_state(raw):
    s = re.sub(r"\s+", " ", str(raw or "").strip().upper())
    if not s or s in AGGREGATE_ROWS:
        return None
    if s in STATE_ALIASES:
        return STATE_ALIASES[s]
    return s.title().replace(" And ", " and ")


def as_int(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = re.sub(r"[^\d\-]", "", str(v))
    return int(s) if s else 0


# ── 1. Read the workbook ─────────────────────────────────────────────────────

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

# Which canonical heads exist in every usable sheet?
per_sheet_heads = []
for sn in SHEETS:
    hdr = [str(h).strip().upper() if h else "" for h in next(wb[sn].iter_rows(values_only=True))]
    per_sheet_heads.append({CANON[h] for h in hdr if h in CANON})
HEADS = sorted(set.intersection(*per_sheet_heads))

# state -> year -> head -> value
data = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
skipped_years = set()

for sn in SHEETS:
    ws = wb[sn]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip().upper() if h else "" for h in next(rows)]
    idx = {}
    for i, h in enumerate(hdr):
        if h in CANON and CANON[h] in HEADS and CANON[h] not in idx:
            idx[CANON[h]] = i
    try:
        iS, iD, iY = hdr.index("STATE"), hdr.index("DISTRICT"), hdr.index("YEAR")
    except ValueError:
        continue

    for r in rows:
        if r is None or iS >= len(r):
            continue
        district = re.sub(r"\s+", " ", str(r[iD] or "").strip().upper())
        if district in AGGREGATE_ROWS:
            continue  # never mix aggregate rows with the districts they aggregate
        st = canon_state(r[iS])
        if not st:
            continue
        m = re.match(r"\s*(\d{4})", str(r[iY] or ""))
        if not m:
            skipped_years.add(sn)
            continue
        y = int(m.group(1))
        if not (2001 <= y <= 2018):
            skipped_years.add(sn)
            continue
        for head, i in idx.items():
            if i < len(r):
                data[st][y][head] += as_int(r[i])

wb.close()

years = sorted({y for st in data for y in data[st]})
states = sorted(data)

# ── 2. Read + simplify the district geometry ─────────────────────────────────


def perp_dist(p, a, b):
    (x, y), (x1, y1), (x2, y2) = p, a, b
    dx, dy = x2 - x1, y2 - y1
    if dx == 0 and dy == 0:
        return math.hypot(x - x1, y - y1)
    t = max(0, min(1, ((x - x1) * dx + (y - y1) * dy) / (dx * dx + dy * dy)))
    return math.hypot(x - (x1 + t * dx), y - (y1 + t * dy))


def rdp(pts, eps):
    """Ramer-Douglas-Peucker, iterative so deep rings can't blow the stack."""
    if len(pts) < 3:
        return pts
    keep = [False] * len(pts)
    keep[0] = keep[-1] = True
    stack = [(0, len(pts) - 1)]
    while stack:
        s, e = stack.pop()
        dmax, imax = 0.0, s
        for i in range(s + 1, e):
            d = perp_dist(pts[i], pts[s], pts[e])
            if d > dmax:
                dmax, imax = d, i
        if dmax > eps:
            keep[imax] = True
            stack.append((s, imax))
            stack.append((imax, e))
    return [p for p, k in zip(pts, keep) if k]


TOL = 0.012  # degrees; ~1.3 km — plenty for a nationwide choropleth
MIN_RING_PTS = 4

geo = json.load(open(GEO, encoding="utf-8"))
shapes = defaultdict(list)  # state -> [path-string, ...]

for feat in geo["features"]:
    st = canon_state(feat["properties"].get("st_nm"))
    if not st:
        continue
    g = feat.get("geometry") or {}
    polys = []
    if g.get("type") == "Polygon":
        polys = [g["coordinates"]]
    elif g.get("type") == "MultiPolygon":
        polys = g["coordinates"]
    for poly in polys:
        for ring in poly:
            raw = [(float(c[0]), float(c[1])) for c in ring if c and len(c) >= 2]
            if len(raw) < MIN_RING_PTS:
                continue
            # Back off the tolerance for features smaller than it — otherwise
            # small island UTs (Lakshadweep) simplify away to nothing and drop
            # off the map entirely.
            pts = []
            for tol in (TOL, TOL / 4, TOL / 16, 0.0):
                pts = rdp(raw, tol) if tol else raw
                if len(pts) >= MIN_RING_PTS:
                    break
            if len(pts) < MIN_RING_PTS:
                continue
            shapes[st].append([[round(x, 3), round(y, 3)] for x, y in pts])

# ── 3. Emit ──────────────────────────────────────────────────────────────────

matched = sorted(set(shapes) & set(states))
data_only = sorted(set(states) - set(shapes))
geo_only = sorted(set(shapes) - set(states))

out_series = {}
for st in states:
    out_series[st] = {
        str(y): {h: data[st][y][h] for h in HEADS if data[st][y][h]} for y in sorted(data[st])
    }

payload = {
    "meta": {
        "years": years,
        "heads": HEADS,
        "labels": LABELS,
        "states_with_data": len(states),
        "states_with_geometry": len(shapes),
        "states_matched": len(matched),
        "data_without_geometry": data_only,
        "geometry_without_data": geo_only,
        "simplify_tolerance_deg": TOL,
        "source": "NCRB district-wise IPC returns, 2001-2015",
        "basis": (
            "Sum of district rows per state-year, excluding the sheets' own aggregate rows. "
            "Runs ~3% above the sheets' TOTAL rows because railway/CID units are included."
        ),
    },
    "shapes": shapes,
    "series": out_series,
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("// GENERATED FILE - do not edit by hand.\n")
    fh.write("// Produced by scripts/build_ncrb_map.py from the NCRB workbook + district GeoJSON.\n")
    fh.write("// Regenerate with:  python3 scripts/build_ncrb_map.py\n\n")
    fh.write("window.INDIA_CRIME_MAP = ")
    json.dump(payload, fh, separators=(",", ":"), ensure_ascii=False)
    fh.write(";\n")

size = os.path.getsize(OUT)
print(f"years        : {years[0]}-{years[-1]} ({len(years)})")
print(f"crime heads  : {', '.join(HEADS)}")
print(f"states data  : {len(states)}   geometry: {len(shapes)}   matched: {len(matched)}")
if data_only:
    print(f"  data but no geometry : {data_only}")
if geo_only:
    print(f"  geometry but no data : {geo_only}")
if skipped_years:
    print(f"sheets skipped for unusable YEAR values: {sorted(skipped_years)}")
tot = sum(data[st][years[-1]]["murder"] for st in states if years[-1] in data[st])
print(f"murders {years[-1]}  : {tot:,}")
print(f"wrote {os.path.relpath(OUT, ROOT)}  ({size/1024:.0f} KB)")
